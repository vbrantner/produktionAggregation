#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import pandas as pd
import calendar
from pathlib import Path
import openpyxl

df = pd.read_csv("~/Documents/coding/bbn/produktionAggregation/data/dataframe.csv")
df['Datum'] = pd.to_datetime(df['Datum'])
df = df.round()

monthList = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
monthList31 = ['Jan', 'Mar', 'May', 'Jul', 'Aug', 'Oct', 'Dec']
monthDict = dict((v,k) for k,v in enumerate(calendar.month_abbr))

def determineYearMonth(df):
    uniqueDates = df.Datum.unique()
    yearList = pd.DatetimeIndex(uniqueDates).year.unique()
    lastMonth = pd.DatetimeIndex(uniqueDates).max().month
    firstMonth = pd.DatetimeIndex(uniqueDates).min().month
    return (yearList, lastMonth, firstMonth)

yearList = determineYearMonth(df)[0]
lastMonth = determineYearMonth(df)[1]
firstMonth = determineYearMonth(df)[2]

def monthToNumber(month):
    if monthDict[month] < 10:
        return '0{}'.format(monthDict[month])
    else:
        return '{}'.format(monthDict[month])

def allMonth_df(df, specificCol):
    if specificCol == 'Artikelgruppe':
        df = df.groupby(['Artikelgruppe', df.Datum.dt.year, df.Datum.dt.month]).sum()
        df = df.GelieferteMenge.unstack()
        return df
    else:
        df = df.groupby(['ArtNr', 'ArtBezeichnung', df.Datum.dt.year, df.Datum.dt.month]).sum()
        df = df.GelieferteMenge.unstack()
        return df

def selectMonth(month, year):
    monthInNumber = monthToNumber(month)
    if month in monthList31:
        return ((df['Datum'] >= '{}-{}-{}'.format(year,monthInNumber,'01')) & (df['Datum'] <= '{}-{}-{}'.format(year, monthInNumber, '31')))
    elif month == "Feb": # todo: SchaltJahre Beruecksichtigen
        return ((df['Datum'] >= '{}-{}-{}'.format(year,monthInNumber,'01')) & (df['Datum'] <= '{}-{}-{}'.format(year, monthInNumber, '28')))
    else:
        return ((df['Datum'] >= '{}-{}-{}'.format(year,monthInNumber,'01')) & (df['Datum'] <= '{}-{}-{}'.format(year, monthInNumber, '30')))

def selectYear(year):
    return ((df['Datum'] >= '{}-{}-{}'.format(year,'01','01')) & (df['Datum'] <= '{}-{}-{}'.format(year, '12', '31')))

def genMonthFilter(year, untilCount, fromMonth):
    string = ""
    for month in monthList[fromMonth : untilCount]:
        string = string + ' | (selectMonth("{}", {}))'.format(month, year)
    return string[3:]

def diffColumn(df, specificCol):
    aggBy = [df.Datum.dt.year, specificCol] 
    if specificCol == 'ArtNr':
        aggBy.append('ArtBezeichnung')
    diff16 = calcDiffHelper(df, 2016, 12, firstMonth, aggBy)
    diff17 = calcDiffHelper(df, 2017, 12, 0, aggBy)
    diff18 = calcDiffHelper(df, 2018, lastMonth, 0, aggBy)
    year18 = df[eval(genMonthFilter(2018, lastMonth, 0))].groupby(aggBy).sum().GelieferteMenge
    year19 = df[eval(genMonthFilter(2019, lastMonth, 0))].groupby(aggBy).sum().GelieferteMenge
    year18 = year18.reset_index('Datum').GelieferteMenge
    year19 = year19.reset_index('Datum').GelieferteMenge
    diff18 = year18 - year19
    diff18 = diff18.to_frame()
    diff18['Datum'] = 2018
    diff19 = year19 - year18
    diff19 = diff19.to_frame()
    diff19['Datum'] = 2019
    result = diff16.append(diff17)
    result = result.append(diff18)
    result = result.append(diff19)
    result = result.set_index('Datum', append=True)
    result = result.rename(columns={'GelieferteMenge': 'Differenz'})
    return result

def calcDiffHelper(df, year, lastMonth, firstMonth, aggBy):
    yearOne = df[eval(genMonthFilter(year, lastMonth, firstMonth))].groupby(aggBy).sum().GelieferteMenge
    yearTwo = df[eval(genMonthFilter(year+1, lastMonth, firstMonth))].groupby(aggBy).sum().GelieferteMenge
    yearOne = yearOne.reset_index('Datum').GelieferteMenge
    yearTwo= yearTwo.reset_index('Datum').GelieferteMenge
    diff = yearTwo - yearOne
    diff = diff.to_frame()
    diff['Datum'] = year
    return diff

def calcYearSum(df, specificCol):
    aggBy = [df.Datum.dt.year, specificCol] 
    if specificCol == 'ArtNr':
        aggBy.append('ArtBezeichnung')
    year16 = df[(selectMonth("Aug", 2016)) | (selectMonth("Oct", 2016)) | (selectMonth("Nov", 2016)) | (selectMonth("Dec", 2016))].groupby(aggBy).sum()
    year17 = df[eval(genMonthFilter(2017, lastMonth, 0))].groupby(aggBy).sum()
    year18 = df[eval(genMonthFilter(2018, lastMonth, 0))].groupby(aggBy).sum()
    year16 = year16.append(year17)
    year16 = year16.append(year18)
    year16 = year16.rename(columns={'GelieferteMenge': 'Summe'})
    year16 = year16.reset_index('Datum')
    year16 = year16.set_index('Datum', append=True)
    return year16.Summe

def createDF(df, specificCol):
    aggDFArt = allMonth_df(df, specificCol)
    diff = diffColumn(df, specificCol)
    yearSum = calcYearSum(df, specificCol)
    aggDFArt = aggDFArt.join(diff, how='outer')
    aggDFArt = aggDFArt.join(yearSum)
    return aggDFArt

def formatDF(df, specificCol):
    df = df[['Summe', 'Differenz',1,2,3,4,5,6,7,8,9,10,11,12]]
    df = df.rename(columns={1:'Jan', 2:'Feb', 3:'Mar', 4:'Apr', 5:'Mai', 6:'Jun', 7:'Jul', 8:'Aug', 9:'Sep', 10:'Oct', 11:'Nov', 12:'Dec'})
    if specificCol == 'Artikelgruppe':
        df = df.loc[['Brote', 'Belegte', 'Brötchen/Kleingebäck', 'Feingebäck', 'Kuchen/Torten',\
        'Snack', 'Teiglinge']]
    elif specificCol == 'ArtNr':
        df = df.drop([12, 13, 27, 49, 19, 25, 26, 28, 29, 31, 32, 34, 35, 39, 41, 42, 43, 44, 46, 47, 48, 60,\
                61, 62, 63, 64, 65, 66, 67, 74, 81, 91, 95, 96, 97, 99, 224, 239, 251, 253, 255, 261, 269, 270,\
                268, 367, 369, 371, 380, 389, 562, 582, 584, 950, 952, 953, 990, 991, 992, 993, 994, 999, 214, 290,\
                1088, 1090, 1091, 1092, 1093, 1094, 1095, 1096, 1097, 1098, 1100])
    df = df.fillna(value="")
    return df

def saveDF():
    artNr = createDF(df, 'ArtNr')
    artNr = formatDF(artNr, 'ArtNr')
    artGrp = createDF(df, 'Artikelgruppe')
    artGrp = formatDF(artGrp, 'Artikelgruppe')

    writer = pd.ExcelWriter('produktionsAnalyse.xlsx', engine='xlsxwriter')
    artGrp.to_excel(writer, 'Artikelgruppen')
    artNr.to_excel(writer, 'ArtikelABC')
    workbook = writer.book
    worksheet = writer.sheets['Artikelgruppen']
    chart1 = workbook.add_chart({'type': 'line'})
    chart2 = workbook.add_chart({'type': 'line'})
    chart3 = workbook.add_chart({'type': 'line'})
    chart4 = workbook.add_chart({'type': 'line'})
    chart5 = workbook.add_chart({'type': 'line'})
    chart6 = workbook.add_chart({'type': 'line'})
    chart7 = workbook.add_chart({'type': 'line'})
    chart1.add_series({
        'name': 'Belegte 17',
        'categories': '=Artikelgruppen!$F$1:$Q$1',
        'values': '=Artikelgruppen!$F$3:$Q$3'})
    chart1.add_series({
        'name': 'Belegete 18',
        'categories': '=Artikelgruppen!$F$1:$Q$1',
        'values': '=Artikelgruppen!$F$4:$P$4'})
    chart2.add_series({
        'name': 'Brote 17',
        'categories': '=Artikelgruppen!$F$1:$Q$1',
        'values': '=Artikelgruppen!$F$7:$Q$7'})
    chart2.add_series({
        'name': 'Brote 18',
        'categories': '=Artikelgruppen!$F$1:$Q$1',
        'values': '=Artikelgruppen!$F$8:$P$8'})
    chart3.add_series({
        'name': 'Brötchen 17',
        'categories': '=Artikelgruppen!$F$1:$Q$1',
        'values': '=Artikelgruppen!$F$11:$Q$11'})
    chart3.add_series({
        'name': 'Brötchen 18',
        'categories': '=Artikelgruppen!$F$1:$Q$1',
        'values': '=Artikelgruppen!$F$12:$P$12'})
    chart4.add_series({
        'name': 'Feingebäck 17',
        'categories': '=Artikelgruppen!$F$1:$Q$1',
        'values': '=Artikelgruppen!$F$15:$Q$15'})
    chart4.add_series({
        'name': 'Feingebäck 18',
        'categories': '=Artikelgruppen!$F$1:$Q$1',
        'values': '=Artikelgruppen!$F$16:$P$16'})
    chart5.add_series({
        'name': 'Kuchen 17',
        'categories': '=Artikelgruppen!$F$1:$Q$1',
        'values': '=Artikelgruppen!$F$19:$Q$19'})
    chart5.add_series({
        'name': 'Kuchen 18',
        'categories': '=Artikelgruppen!$F$1:$Q$1',
        'values': '=Artikelgruppen!$F$20:$P$20'})
    chart6.add_series({
        'name': 'Snack 17',
        'categories': '=Artikelgruppen!$F$1:$Q$1',
        'values': '=Artikelgruppen!$F$23:$Q$23'})
    chart6.add_series({
        'name': 'Snack 18',
        'categories': '=Artikelgruppen!$F$1:$Q$1',
        'values': '=Artikelgruppen!$F$24:$P$24'})
    chart7.add_series({
        'name': 'Teiglinge 17',
        'categories': '=Artikelgruppen!$F$1:$Q$1',
        'values': '=Artikelgruppen!$F$27:$Q$27'})
    chart7.add_series({
        'name': 'Teiglinge 18',
        'categories': '=Artikelgruppen!$F$1:$Q$1',
        'values': '=Artikelgruppen!$F$28:$P$28'})
    chart1.set_size({'x_scale': 1.4, 'y_scale': 1.0})
    chart2.set_size({'x_scale': 1.4, 'y_scale': 1.0})
    chart3.set_size({'x_scale': 1.4, 'y_scale': 1.0})
    chart4.set_size({'x_scale': 1.4, 'y_scale': 1.0})
    chart5.set_size({'x_scale': 1.4, 'y_scale': 1.0})
    chart6.set_size({'x_scale': 1.4, 'y_scale': 1.0})
    chart7.set_size({'x_scale': 1.4, 'y_scale': 1.0})
    worksheet.insert_chart('A30', chart1)
    worksheet.insert_chart('A45', chart2)
    worksheet.insert_chart('A60', chart3)
    worksheet.insert_chart('A75', chart4)
    worksheet.insert_chart('L30', chart5)
    worksheet.insert_chart('L45', chart6)
    worksheet.insert_chart('L60', chart7)

    writer.save()

if __name__ == "__main__":
    # execute only if run as a script
    saveDF()



